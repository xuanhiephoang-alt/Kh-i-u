"""Bang tag chung -> kiem tra dia chi -> xuat cho GX Works3 (Mitsubishi), KV STUDIO (Keyence), EasyBuilder Pro (Weintek).

Bang tag (CSV/Excel), cot:
    tag      - ten tag/nhan (bat buoc, khong trung, khong dau cach)
    dia_chi  - dia chi PLC (bat buoc): X10, Y0, M100, D200, D200.5, R1015, MR500, DM100...
    kieu     - BOOL / INT / UINT / WORD / DINT / UDINT / DWORD / REAL (bo trong: BOOL voi thiet bi bit, INT voi thanh ghi)
    mo_ta    - chu thich
    plc      - ten PLC (khi du an co nhieu PLC). Bo trong = PLC duy nhat khai bao bang --plc
    hmi      - x/1/co: dua tag nay sang HMI. Bo trong ca cot = dua tat ca
    rw       - R hoac RW (cho HMI, mac dinh RW)

Dinh dang file import cua phan mem thay doi theo phien ban/ngon ngu. Dau ra mac dinh la dinh dang
pho bien nhat nhung CHUA kiem chung tren tung phien ban: hay xuat 1 file mau tu phan mem cua ban va
truyen vao --template-* de cong cu khop dung cot, dau phan cach va encoding.
"""
from __future__ import annotations

import argparse
import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from .pricelist import norm

# ------------------------------------------------------------------ dong PLC

# (tien to, loai, he co so). Xep tien to dai truoc de "SM" khong bi doc thanh "S"+"M".
_MITSU_COMMON = [
    ("SM", "bit", 10), ("SB", "bit", 16), ("SD", "word", 10), ("SW", "word", 16), ("ZR", "word", 10),
    ("M", "bit", 10), ("L", "bit", 10), ("F", "bit", 10), ("B", "bit", 16),
    ("D", "word", 10), ("W", "word", 16), ("R", "word", 10), ("T", "any", 10), ("C", "any", 10),
]
FAMILIES = {
    # FX5/FX3: X/Y danh so bat phan (X0-X7, X10...)
    "fx5": [("X", "bit", 8), ("Y", "bit", 8)] + _MITSU_COMMON,
    # iQ-R / Q / L: X/Y he 16
    "iqr": [("X", "bit", 16), ("Y", "bit", 16)] + _MITSU_COMMON,
    # Keyence KV-8000/7000/5000/NANO: R/MR/LR/CR = kenh*100 + bit(00-15)
    "kv": [("MR", "chbit", 10), ("LR", "chbit", 10), ("CR", "chbit", 10), ("R", "chbit", 10),
           ("DM", "word", 10), ("EM", "word", 10), ("FM", "word", 10), ("ZF", "word", 10),
           ("TM", "word", 10), ("CM", "word", 10), ("VB", "bit", 16), ("VM", "word", 10),
           ("B", "bit", 16), ("W", "word", 16), ("T", "any", 10), ("C", "any", 10)],
}
FAMILIES["q"] = FAMILIES["iqr"]
VENDOR = {"fx5": "mitsubishi", "iqr": "mitsubishi", "q": "mitsubishi", "kv": "keyence"}

WORD_TYPES = {"INT": 1, "UINT": 1, "WORD": 1, "DINT": 2, "UDINT": 2, "DWORD": 2, "REAL": 2}
ALL_TYPES = {"BOOL", *WORD_TYPES}
_DIGITS = {8: "01234567", 10: "0123456789", 16: "0123456789ABCDEF"}


@dataclass
class Tag:
    tag: str
    address: str
    dtype: str = ""
    comment: str = ""
    plc: str = ""
    hmi: bool = True
    rw: str = "RW"
    row: int = 0
    # dien sau khi phan tich
    prefix: str = ""
    number: str = ""
    bit: int | None = None
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def parse_address(addr: str, family: str):
    """'D200.5' -> ('D', '200', 5, kind, radix). Raise ValueError neu sai."""
    a = addr.strip().upper().replace(" ", "")
    for prefix, kind, radix in FAMILIES[family]:
        if not a.startswith(prefix):
            continue
        rest = a[len(prefix):]
        num, _, bit = rest.partition(".")
        # "DM100" voi FX5: "D" + "M100" -> khong phai tien to nay, thu tien to khac
        if not num or (radix != 16 and not num[0].isdigit()) or num[0] not in _DIGITS[16]:
            continue
        if any(ch not in _DIGITS[radix] for ch in num):
            base = {8: "bat phan (0-7)", 10: "thap phan", 16: "thap luc phan"}[radix]
            raise ValueError(f"{prefix} danh so {base}: '{num}' khong hop le")
        if kind == "chbit" and int(num) % 100 > 15:
            raise ValueError(f"{prefix}{num}: 2 so cuoi la bit 00-15 (vd {prefix}{int(num) // 100}15)")
        bit_i = None
        if bit:
            if kind != "word":
                raise ValueError(f"{prefix} la thiet bi bit, khong dung '.bit'")
            try:
                bit_i = int(bit, 16) if family != "kv" else int(bit)
            except ValueError:
                raise ValueError(f"bit '{bit}' khong hop le") from None
            if not 0 <= bit_i <= 15:
                raise ValueError(f"bit {bit} ngoai khoang 0-15")
        return prefix, num, bit_i, kind, radix
    known = ", ".join(p for p, _, _ in FAMILIES[family])
    raise ValueError(f"tien to khong ho tro voi {family} (dung: {known})")


_TAG_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def validate(tags: list[Tag], families: dict[str, str]) -> None:
    """Ghi loi/canh bao vao tung tag: dinh dang dia chi, kieu du lieu, trung ten, chong lan vung nho."""
    names: dict[str, Tag] = {}
    words: dict[tuple, Tag] = {}   # (plc, prefix, word_index) -> tag dung word
    bits: dict[tuple, Tag] = {}    # (plc, address chuan) -> tag
    for t in tags:
        fam = families.get(t.plc)
        if fam is None:
            t.errors.append(f"PLC '{t.plc}' chua khai bao dong (--plc {t.plc}=fx5|iqr|q|kv)")
            continue
        if not _TAG_RE.match(t.tag):
            t.errors.append("ten tag chi gom chu khong dau, so, '_' va khong bat dau bang so")
        key = (t.plc, t.tag.upper())
        if key in names:
            t.errors.append(f"trung ten tag voi dong {names[key].row}")
        names[key] = t
        try:
            t.prefix, t.number, t.bit, kind, radix = parse_address(t.address, fam)
        except ValueError as e:
            t.errors.append(str(e))
            continue
        is_bit = kind in ("bit", "chbit") or t.bit is not None
        if not t.dtype:
            t.dtype = "BOOL" if is_bit or kind == "any" else "INT"
        if t.dtype not in ALL_TYPES:
            t.errors.append(f"kieu '{t.dtype}' khong ho tro ({', '.join(sorted(ALL_TYPES))})")
            continue
        if is_bit and t.dtype != "BOOL":
            t.errors.append(f"{t.address} la bit, kieu phai la BOOL")
            continue
        if kind == "word" and t.bit is None and t.dtype == "BOOL":
            t.errors.append(f"{t.address} la thanh ghi word; dung kieu so hoac dia chi bit dang {t.prefix}{t.number}.0")
            continue
        if t.dtype in WORD_TYPES and kind in ("word", "any"):
            n = int(t.number, radix)
            for i in range(WORD_TYPES[t.dtype]):
                k = (t.plc, t.prefix, n + i)
                if k in words and words[k] is not t:
                    o = words[k]
                    t.errors.append(f"chong lan vung nho voi {o.tag} ({o.address} {o.dtype}, dong {o.row})")
                    break
                words[k] = t
        else:
            k = (t.plc, t.prefix + t.number + ("" if t.bit is None else f".{t.bit}"))
            if k in bits:
                t.warnings.append(f"cung dia chi voi {bits[k].tag} (dong {bits[k].row}) - kiem tra co co y khong")
            bits[k] = t
    for t in tags:
        if t.errors or not t.prefix:
            continue
        fam = families[t.plc]
        n = int(t.number, next(r for p, _, r in FAMILIES[fam] if p == t.prefix))
        # bit-cua-word nam trong mot word da khai bao kieu so
        if t.bit is not None and (o := words.get((t.plc, t.prefix, n))):
            t.warnings.append(f"la bit cua {o.tag} ({o.address})")
        if WORD_TYPES.get(t.dtype) == 2 and n % 2:
            t.warnings.append("du lieu 32-bit o dia chi le - nen dat dia chi chan cho de quan ly")


# ------------------------------------------------------------------ doc bang tag

def read_tags(path: str | Path, default_plc: str) -> list[Tag]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        ws = openpyxl.load_workbook(path, data_only=True).active
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    head = [norm(h).replace(" ", "_") for h in rows[0]]
    for need in ("tag", "dia_chi"):
        if need not in head:
            raise ValueError(f"Bang tag thieu cot '{need}' (co: {head})")
    has_hmi = "hmi" in head
    tags = []
    for n, r in enumerate(rows[1:], 2):
        d = {k: (r[i].strip() if i < len(r) else "") for i, k in enumerate(head)}
        if not d["tag"] and not d["dia_chi"]:
            continue
        tags.append(Tag(
            tag=d["tag"], address=d["dia_chi"].upper().replace(" ", ""), dtype=d.get("kieu", "").upper(),
            comment=d.get("mo_ta", ""), plc=d.get("plc") or default_plc,
            hmi=(norm(d.get("hmi", "")) in ("x", "1", "co", "y", "yes", "true")) if has_hmi else True,
            rw="R" if d.get("rw", "").upper() == "R" else "RW", row=n))
    return tags


# ------------------------------------------------------------------ du lieu cho tung phan mem

GX_TYPE = {"BOOL": "Bit", "INT": "Word [Signed]", "UINT": "Word [Unsigned]/Bit String [16-bit]",
           "WORD": "Word [Unsigned]/Bit String [16-bit]", "DINT": "Double Word [Signed]",
           "UDINT": "Double Word [Unsigned]/Bit String [32-bit]",
           "DWORD": "Double Word [Unsigned]/Bit String [32-bit]", "REAL": "FLOAT [Single Precision]"}
WT_TYPE = {"BOOL": "Bit", "INT": "16-bit Signed", "UINT": "16-bit Unsigned", "WORD": "16-bit Unsigned",
           "DINT": "32-bit Signed", "UDINT": "32-bit Unsigned", "DWORD": "32-bit Unsigned", "REAL": "32-bit Float"}


def fields_for(t: Tag, target: str) -> dict[str, str]:
    """Gia tri cac truong de dien vao cot file xuat."""
    full = t.prefix + t.number + ("" if t.bit is None else f".{t.bit:X}" if target != "kv" else f".{t.bit}")
    return {
        "tag": t.tag, "comment": t.comment, "plc": t.plc, "class": "VAR_GLOBAL",
        "rw": "Read" if t.rw == "R" else "Read/Write",
        "type": WT_TYPE[t.dtype] if target == "weintek" else GX_TYPE[t.dtype] if target == "gx3" else t.dtype,
        # Weintek tach "Address type" (tien to) va "Address" (so); phan mem khac dung dia chi day du
        "prefix": t.prefix,
        "address": (t.number + ("" if t.bit is None else f".{t.bit}")) if target == "weintek" else full,
    }


# Nhan dien cot trong file mau (sau norm()). Thu tu quan trong: "device name" truoc "device".
COLUMN_KEYS = [
    ("device name", "plc"), ("plc", "plc"),
    ("address type", "prefix"), ("loai dia chi", "prefix"),
    ("data format", "type"), ("data type", "type"), ("データ型", "type"), ("kieu", "type"),
    ("read/write", "rw"), ("クラス", "class"), ("class", "class"),
    # "Assign (Device/Label)" chua chu "label": phai xet truoc "label"
    ("assign", "address"), ("割付", "address"),
    ("label", "tag"), ("ラベル", "tag"), ("tag", "tag"),
    ("device", "address"), ("デバイス", "address"),
    ("address", "address"), ("dia chi", "address"),
    ("comment", "comment"), ("コメント", "comment"), ("description", "comment"), ("mo ta", "comment"),
]


def column_field(header: str) -> str | None:
    h = norm(header).replace("　", "")
    for key, f in COLUMN_KEYS:
        if norm(key) in h:  # norm() ca hai phia: no bo dau ゛ cua tieng Nhat
            return f
    return None


DEFAULT_HEADERS = {
    # Weintek EasyBuilder Pro - Address Tag Library > Import CSV
    "weintek": ["Tag name", "Device name", "Address type", "Address", "Data format", "Read/Write", "Comment"],
    # GX Works3 - Global label (ban tieng Anh)
    "gx3": ["Label Name", "Data Type", "Class", "Assign (Device/Label)", "Comment"],
    # Chu thich thiet bi (GX Works3 / KV STUDIO)
    "comment": ["Device", "Comment"],
}


@dataclass
class Template:
    header: list[str]
    preamble: list[str] = field(default_factory=list)
    delimiter: str = ","
    encoding: str = "utf-8-sig"
    newline: str = "\r\n"


def read_template(path: str | Path) -> Template:
    """Doc file mau do phan mem xuat ra: giu nguyen dong dau, cot, dau phan cach, encoding."""
    raw = Path(path).read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        enc = "utf-16"
    elif raw[:3] == b"\xef\xbb\xbf":
        enc = "utf-8-sig"
    else:
        try:
            raw.decode("utf-8")
            enc = "utf-8"
        except UnicodeDecodeError:
            enc = "cp932"  # ban tieng Nhat
    text = raw.decode(enc)
    delim = "\t" if text.count("\t") > text.count(",") else ","
    lines = text.splitlines()
    # dong tieu de = dong dau tien co >= 1 cot nhan dien duoc va nhieu cot nhat
    best, best_score = 0, -1
    for i, line in enumerate(lines[:20]):
        cols = next(csv.reader([line], delimiter=delim), [])
        score = sum(column_field(c) is not None for c in cols)
        if score > best_score:
            best, best_score = i, score
    header = next(csv.reader([lines[best]], delimiter=delim))
    return Template(header, lines[:best], delim, "utf-16" if enc == "utf-16" else enc,
                    "\r\n" if "\r\n" in text else "\n")


def write_table(path: str | Path, tags: list[Tag], target: str, tpl: Template) -> int:
    cols = [column_field(h) for h in tpl.header]
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=tpl.delimiter, lineterminator=tpl.newline, quoting=csv.QUOTE_ALL)
    for line in tpl.preamble:
        buf.write(line + tpl.newline)
    w.writerow(tpl.header)
    for t in tags:
        f = fields_for(t, target)
        w.writerow([f.get(c, "") if c else "" for c in cols])
    Path(path).write_bytes(buf.getvalue().encode(tpl.encoding))
    return len(tags)


# ------------------------------------------------------------------ bao cao Excel

def write_report(path: str | Path, tags: list[Tag], families: dict[str, str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Tag"
    ws.append(["Dong", "PLC", "Dong PLC", "Tag", "Dia chi", "Kieu", "HMI", "R/W", "Mo ta", "Loi", "Canh bao"])
    red, yellow = PatternFill("solid", fgColor="F8CBAD"), PatternFill("solid", fgColor="FFF2CC")
    for t in tags:
        ws.append([t.row, t.plc, families.get(t.plc, "?"), t.tag, t.address, t.dtype, "x" if t.hmi else "",
                   t.rw, t.comment, "; ".join(t.errors), "; ".join(t.warnings)])
        fill = red if t.errors else yellow if t.warnings else None
        if fill:
            for c in ws[ws.max_row]:
                c.fill = fill
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJK", [6, 10, 8, 22, 10, 8, 5, 5, 34, 50, 40]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# ------------------------------------------------------------------ CLI

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="python -m dienkit.tags",
                                 description="Bang tag chung -> GX Works3 / KV STUDIO / Weintek EasyBuilder Pro")
    ap.add_argument("tags", help="Bang tag (.csv/.xlsx)")
    ap.add_argument("--plc", action="append", required=True, metavar="TEN=DONG",
                    help="Khai bao PLC: --plc PLC1=fx5 (fx5 | iqr | q | kv). Lap lai neu nhieu PLC. "
                         "PLC dau tien la mac dinh cho dong de trong cot plc. TEN = ten thiet bi trong EasyBuilder.")
    ap.add_argument("-o", "--out", default="tags_out", help="Thu muc xuat")
    ap.add_argument("--template-weintek", help="File CSV xuat tu EasyBuilder (Address Tag Library > Export)")
    ap.add_argument("--template-gx3", help="File CSV nhan toan cuc xuat tu GX Works3")
    ap.add_argument("--template-comment", help="File CSV chu thich thiet bi xuat tu GX Works3/KV STUDIO")
    ap.add_argument("--force", action="store_true", help="Van xuat file khi con loi")
    a = ap.parse_args(argv)

    families: dict[str, str] = {}
    for spec in a.plc:
        name, _, fam = spec.partition("=")
        if fam not in FAMILIES:
            ap.error(f"--plc {spec}: dong PLC phai la fx5, iqr, q hoac kv")
        families[name] = fam
    default_plc = next(iter(families))

    tags = read_tags(a.tags, default_plc)
    validate(tags, families)
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    write_report(out / "kiem_tra_tag.xlsx", tags, families)

    n_err = sum(bool(t.errors) for t in tags)
    for t in tags:
        for e in t.errors:
            print(f"LOI    dong {t.row} {t.tag} {t.address}: {e}")
        for w in t.warnings:
            print(f"CANH BAO dong {t.row} {t.tag} {t.address}: {w}")
    print(f"{len(tags)} tag, {n_err} tag loi. Bao cao: {out / 'kiem_tra_tag.xlsx'}")
    if n_err and not a.force:
        print("Chua xuat file import vi con loi (sua bang tag hoac dung --force).")
        return 1

    ok = [t for t in tags if not t.errors]
    tpl = lambda p, kind: read_template(p) if p else Template(DEFAULT_HEADERS[kind])  # noqa: E731
    for plc, fam in families.items():
        mine = [t for t in ok if t.plc == plc]
        if not mine:
            continue
        cm = [t for t in mine if t.comment]
        if VENDOR[fam] == "mitsubishi":
            n = write_table(out / f"{plc}_gx3_global_label.csv", mine, "gx3", tpl(a.template_gx3, "gx3"))
            print(f"  {plc}: {out / f'{plc}_gx3_global_label.csv'} ({n} nhan)")
            n = write_table(out / f"{plc}_gx3_device_comment.csv", cm, "gx3", tpl(a.template_comment, "comment"))
            print(f"  {plc}: {out / f'{plc}_gx3_device_comment.csv'} ({n} chu thich)")
        else:
            n = write_table(out / f"{plc}_kv_device_comment.csv", cm, "kv", tpl(a.template_comment, "comment"))
            print(f"  {plc}: {out / f'{plc}_kv_device_comment.csv'} ({n} chu thich)")
    hmi = [t for t in ok if t.hmi]
    if any(t.bit is not None for t in hmi):
        print("  LUU Y: tag dang bit-cua-word (vd D30.0) - kiem tra lai cach EasyBuilder nhap dia chi nay sau khi import.")
    n = write_table(out / "weintek_address_tags.csv", hmi, "weintek", tpl(a.template_weintek, "weintek"))
    print(f"  HMI: {out / 'weintek_address_tags.csv'} ({n} tag)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
