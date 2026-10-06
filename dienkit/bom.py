"""Dung bang vat tu (BOM) tu ket qua tinh tai + bang gia, xuat Excel co cong thuc."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from .calc import Load, Result, size_load, size_main
from .pricelist import Item, search


@dataclass
class BomLine:
    group: str
    query: str
    qty: float
    unit: str
    item: Item | None
    score: float


def read_loads(path: str | Path) -> list[Load]:
    """Doc danh sach tai tu CSV hoac Excel. Cot: ten,kw,pha,dien_ap,cosphi,hieu_suat,chieu_dai_m,kd,dong_co,so_loi"""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        ws = openpyxl.load_workbook(path, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    head = [str(h or "").strip().lower() for h in rows[0]]
    loads = []
    for r in rows[1:]:
        if not r or not str(r[0] or "").strip():
            continue
        d = dict(zip(head, r))

        def num(k, default):
            v = d.get(k)
            return default if v in (None, "") else float(str(v).replace(",", "."))

        loads.append(Load(
            name=str(d["ten"]).strip(), kw=num("kw", 0), phases=int(num("pha", 3)),
            voltage=num("dien_ap", 380 if int(num("pha", 3)) == 3 else 220),
            pf=num("cosphi", 0.85), eff=num("hieu_suat", 1.0), length_m=num("chieu_dai_m", 20),
            kd=num("kd", 1.0), motor=str(d.get("dong_co") or "").strip().lower() in ("1", "x", "y", "yes", "co", "có", "true"),
            cores=int(num("so_loi", 0))))
    return loads


def build_bom(results: list[Result], items: list[Item], cable_prefix: str = "CVV") -> list[BomLine]:
    lines: list[BomLine] = []
    for r in results:
        if r.breaker_in:
            kind = "MCCB" if r.breaker_in > 63 else "MCB"
            q = f"{kind} {r.breaker_poles}P {r.breaker_in}A"
            it, sc = search(items, q)
            lines.append(BomLine(r.load.name, q, 1, "cai", it, sc))
        if r.section:
            cores = r.load.cores or (4 if r.load.phases == 3 else 3)
            s = int(r.section) if float(r.section).is_integer() else r.section
            q = f"{cable_prefix} {cores}x{s}"
            it, sc = search(items, q)
            lines.append(BomLine(r.load.name, q, r.load.length_m, "m", it, sc))
    return lines


def write_excel(path: str | Path, results: list[Result], main: dict, lines: list[BomLine],
                vat: float = 0.10) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    warn = PatternFill("solid", fgColor="FFF2CC")
    bad = PatternFill("solid", fgColor="F8CBAD")

    wb = Workbook()
    ws = wb.active
    ws.title = "Tinh tai"
    ws.append(["Tai", "kW", "Pha", "U (V)", "cosphi", "Kd", "kVA", "Ib (A)", "Aptomat In (A)",
               "Cuc", "Tiet dien (mm2)", "Iz (A)", "Dai (m)", "Sut ap (%)", "Ghi chu"])
    for r in results:
        L = r.load
        ws.append([L.name, L.kw, L.phases, L.voltage, L.pf, L.kd, round(r.kva, 2), round(r.ib, 1),
                   r.breaker_in, r.breaker_poles, r.section, round(r.iz, 1) if r.iz else None,
                   L.length_m, round(r.vdrop_pct, 2) if r.vdrop_pct is not None else None, r.note])
    n = len(results) + 2
    ws.append([])
    ws.append(["Tong (Ks=%.2f)" % main["ks"], round(main["kw"], 2), None, None, None, None,
               round(main["kva"], 2), round(main["i"], 1), main["breaker"]])
    ws.cell(row=n + 1, column=1).font = bold
    widths = [24, 8, 6, 8, 8, 6, 9, 9, 14, 6, 14, 9, 8, 10, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    wb2 = wb.create_sheet("Vat tu")
    wb2.append(["STT", "Hang muc", "Yeu cau", "Mat hang trong bang gia", "Nguon", "DVT", "SL", "Don gia", "Thanh tien", "Do khop"])
    for i, ln in enumerate(lines, 1):
        row = i + 1
        it = ln.item
        wb2.append([i, ln.group, ln.query, it.name if it else "(KHONG TIM THAY - nhap tay)",
                    it.source if it else "", ln.unit, ln.qty, it.price if it else None,
                    f"=IF(H{row}=\"\",0,G{row}*H{row})", round(ln.score, 2)])
        fill = bad if not it else (warn if ln.score < 1 else None)
        if fill:
            for c in wb2[row]:
                c.fill = fill
    last = len(lines) + 1
    wb2.append([])
    wb2.append([None, "Cong chua VAT", None, None, None, None, None, None, f"=SUM(I2:I{last})"])
    wb2.append([None, f"VAT {vat:.0%}", None, None, None, None, None, None, f"=I{last + 2}*{vat}"])
    wb2.append([None, "Tong cong", None, None, None, None, None, None, f"=I{last + 2}+I{last + 3}"])
    for r in (last + 2, last + 3, last + 4):
        wb2.cell(row=r, column=2).font = bold
        wb2.cell(row=r, column=9).font = bold
    for col in "HI":
        for c in wb2[col][1:]:
            c.number_format = "#,##0"
    for i, w in enumerate([5, 22, 22, 46, 24, 6, 8, 14, 16, 8], 1):
        wb2.column_dimensions[chr(64 + i)].width = w

    for sheet in (ws, wb2):
        for c in sheet[1]:
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(horizontal="center", wrap_text=True)
    wb2.cell(row=last + 6, column=2, value="Vang = khop chua day du, can xac nhan. Do = khong co trong bang gia.")
    wb.save(path)
