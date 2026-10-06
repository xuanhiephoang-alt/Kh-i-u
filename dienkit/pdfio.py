"""Doc ban ve dien dang PDF (xuat tu CAD, co lop chu - khong phai ban scan):
- danh sach I/O PLC: dia chi X/Y + mo ta nam canh nhau tren cung dong
- danh sach ma thiet bi trong ngoac, vd (S8VK-C48024) -> BOM nhap

Chi doc chu nam ngang. Mo ta xuong 2 dong duoc ghep lai. Ban ve scan (anh) can OCR, khong ho tro.
"""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

ADDR = re.compile(r"^[XY][0-7]+$")
# nhan khong phai mo ta: tham chieu cheo (E031-A3), bus nguon, so domino
NOT_DESC = re.compile(r"^\(?E\d{3}-[A-F]\d\)?$|^[PN]24[A-Z]?$|^(TB|XT)\d")
DRAWING_CODE = re.compile(r"^[A-Z0-9]+-([A-Z]\d{3})$")
# ma thiet bi: co ca chu va so, khong phai tham chieu cheo
MODEL = re.compile(r"\(([A-Z0-9][A-Za-z0-9\-/ .]{3,50})\)")
CROSS_REF = re.compile(r"^E\d{3}-[A-F]\d$")


@dataclass
class DrawingIO:
    address: str
    desc: str
    sheet: str   # ma to ban ve, vd E300
    page: int


def _mid(w) -> float:
    return (w["top"] + w["bottom"]) / 2


def extract_io(path: str | Path) -> dict[str, DrawingIO]:
    """Tra ve {dia chi: DrawingIO}. Neu 1 dia chi co mo ta o nhieu to, giu to dau tien."""
    import pdfplumber
    out: dict[str, DrawingIO] = {}
    with pdfplumber.open(path) as pdf:
        for n, page in enumerate(pdf.pages, 1):
            ws = [w for w in page.extract_words(keep_blank_chars=True, x_tolerance=1.5) if w["upright"]]
            sheet = next((m.group(1) for w in ws if (m := DRAWING_CODE.match(w["text"].strip()))), "")
            for a in ws:
                addr = a["text"].strip()
                if not ADDR.match(addr):
                    continue
                row = [w for w in ws if w is not a and abs(_mid(w) - _mid(a)) < 4
                       and not ADDR.match(w["text"].strip()) and not NOT_DESC.match(w["text"].strip())]
                right = [w for w in row if 0 < w["x0"] - a["x1"] < 60]
                left = [w for w in row if 0 < a["x0"] - w["x1"] < 100]
                if not (right or left):
                    continue  # nhan so day, khong co mo ta ben canh
                w = min(right or left, key=lambda w: min(abs(w["x0"] - a["x1"]), abs(a["x0"] - w["x1"])))
                # mo ta ben trai thuong can le phai; ben phai can le trai
                edge = "x1" if not right else "x0"
                desc = w["text"].strip()
                for v in ws:
                    if abs(v[edge] - w[edge]) < 3 and not ADDR.match(v["text"].strip()):
                        if 0 < v["top"] - w["bottom"] < 4:
                            desc = desc + " " + v["text"].strip()
                        elif 0 < w["top"] - v["bottom"] < 4:
                            desc = v["text"].strip() + " " + desc
                if addr not in out:
                    out[addr] = DrawingIO(addr, desc, sheet, n)
    return out


@dataclass
class ModelRef:
    model: str
    count: int = 0
    pages: list[int] = field(default_factory=list)


def extract_models(path: str | Path) -> list[ModelRef]:
    """Ma thiet bi ghi trong ngoac tren ban ve. count = so lan xuat hien (KHONG phai so luong mua)."""
    import pdfplumber
    found: dict[str, ModelRef] = {}
    with pdfplumber.open(path) as pdf:
        for n, page in enumerate(pdf.pages, 1):
            for m in MODEL.finditer(page.extract_text() or ""):
                # "(NV63-SV 2P 30A 100-240V 30mA)" -> NV63-SV ; "(AX9000TS -U0)" -> AX9000TS-U0
                code = re.sub(r"\s+(?=[-/])", "", m.group(1).strip()).split()[0]
                # ma hang thuong co dau '-' (LIG1, CAM... la ky hieu thiet bi, khong phai ma)
                if CROSS_REF.match(code) or "-" not in code or not (re.search(r"\d", code) and re.search(r"[A-Z]", code)):
                    continue
                ref = found.setdefault(code, ModelRef(code))
                ref.count += 1
                if n not in ref.pages:
                    ref.pages.append(n)
    return sorted(found.values(), key=lambda r: r.model)


def write_io_csv(path: str | Path, io: dict[str, DrawingIO], module: str = "") -> None:
    """Ghi theo dinh dang bang I/O cua dienkit.iodraw."""
    def key(a):
        return a[0], int(a[1:], 8)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["dia_chi", "loai", "tag", "mo_ta", "thiet_bi", "dau_day", "module", "to_ban_ve"])
        for a in sorted(io, key=key):
            d = io[a]
            w.writerow([a, "DI" if a[0] == "X" else "DO", "", d.desc, "", "", module, d.sheet])


def write_models_xlsx(path, models: list[ModelRef], items) -> None:
    """BOM nhap: ma | so lan xuat hien | trang | SL (nguoi dung dien) | mat hang trong bang gia | don gia | thanh tien."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from .pricelist import find_by_code
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM nhap"
    ws.append(["STT", "Ma thiet bi", "So lan xuat hien", "Trang ban ve", "SL (nhap tay)",
               "Mat hang trong bang gia", "Nguon", "Don gia", "Thanh tien"])
    red = PatternFill("solid", fgColor="F8CBAD")
    for i, m in enumerate(models, 1):
        r = i + 1
        it = find_by_code(items, m.model) if items else None
        ws.append([i, m.model, m.count, " ".join(map(str, m.pages)), None,
                   it.name if it else ("(khong co trong bang gia)" if items else ""), it.source if it else "",
                   it.price if it else None, f"=IF(OR(E{r}=\"\",H{r}=\"\"),0,E{r}*H{r})"])
        if items and not it:
            for c in ws[r]:
                c.fill = red
    last = len(models) + 1
    ws.append([])
    ws.append([None, "Cong", None, None, None, None, None, None, f"=SUM(I2:I{last})"])
    ws.append([None, "Luu y: so lan xuat hien tren ban ve KHONG phai so luong mua - nhap cot SL."])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in "HI":
        for c in ws[col][1:]:
            c.number_format = "#,##0"
    for col, w in zip("ABCDEFGHI", [5, 24, 10, 12, 10, 40, 22, 14, 16]):
        ws.column_dimensions[col].width = w
    wb.save(path)


def main(argv=None) -> int:
    import argparse
    ap = argparse.ArgumentParser(prog="python -m dienkit.pdfio",
                                 description="Doc danh sach I/O va ma thiet bi tu ban ve PDF")
    ap.add_argument("pdf")
    ap.add_argument("--io", help="Ghi danh sach I/O ra CSV (dinh dang cua dienkit.iodraw)")
    ap.add_argument("--models", help="Ghi danh sach ma thiet bi ra .csv hoac .xlsx (BOM nhap)")
    ap.add_argument("-p", "--prices", nargs="*", default=[], help="Bang gia (.xlsx/.csv/.pdf) de tra gia theo ma")
    a = ap.parse_args(argv)
    io = extract_io(a.pdf)
    by_sheet = defaultdict(int)
    for d in io.values():
        by_sheet[d.sheet] += 1
    print(f"I/O doc duoc: {len(io)} diem ({', '.join(f'{s}: {n}' for s, n in sorted(by_sheet.items()))})")
    if a.io:
        write_io_csv(a.io, io)
        print(f"Da ghi {a.io}")
    models = extract_models(a.pdf)
    print(f"Ma thiet bi: {len(models)}")
    if a.models:
        items = []
        if a.prices:
            from .pricelist import load_many
            items = load_many(a.prices)
        if a.models.lower().endswith(".xlsx"):
            write_models_xlsx(a.models, models, items)
        else:
            with open(a.models, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(["ma", "so_lan_xuat_hien", "trang", "so_luong"])
                for m in models:
                    w.writerow([m.model, m.count, " ".join(map(str, m.pages)), ""])
        print(f"Da ghi {a.models} (cot so luong de trong: so lan xuat hien KHONG phai so luong mua)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
