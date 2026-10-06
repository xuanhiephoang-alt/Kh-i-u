"""Doc project GX Works3 (.gx3, dong FX5) - CHI DOC, khong sua file:
- chu thich thiet bi (device comment)
- thiet bi duoc dung trong chuong trinh ladder (tham chieu cheo)
- model CPU / module mo rong

Va xuat:
- bang I/O cho dienkit.iodraw, bang tag cho dienkit.tags
- bao cao Excel: thiet bi dung ma khong chu thich, chu thich ma khong dung,
  va so sanh chu thich chuong trinh voi ban ve PDF (dienkit.pdfio)

CANH BAO: .gx3 la dinh dang noi bo cua Mitsubishi, khong co tai lieu cong khai. Phan doc o day duoc
suy ra tu file thuc te (GX Works3, FX5U) va co the sai voi phien ban khac. Moi ket qua phai doi chieu
lai trong GX Works3.
"""
from __future__ import annotations

import csv
import difflib
import re
import sqlite3
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path

from .pricelist import norm

# DevCode trong bang DEVICE_DATA (file *_DC.db) -> ten thiet bi. Suy ra bang cach doi chieu chu thich.
DEV_CODES = {1: "M", 2: "SM", 3: "L", 16: "X", 17: "Y", 32: "D", 33: "SD", 66: "T"}
SYSTEM = {"SM", "SD"}
OCTAL = {"X", "Y"}  # FX5: X/Y hien thi bat phan
LADDER_DEVICES = {"X", "Y", "M", "L", "F", "B", "SM", "SB", "D", "W", "R", "ZR", "SD", "SW",
                  "T", "ST", "C", "LC", "LT", "LST", "Z", "LZ", "V"}


def fmt(dev: str, n: int, bit: int = 0) -> str:
    s = f"{dev}{n:o}" if dev in OCTAL else f"{dev}{n}"
    return s + (f".{bit:X}" if bit else "")


@dataclass
class Project:
    comments: dict[str, str]     # "X10" -> "WORK INPUT"
    used: set[str]               # thiet bi xuat hien trong ladder
    models: list[str]            # ["FX5U-80MT/ES", "FX5-32ER/ES"]
    unknown_codes: dict[int, int]  # DevCode chua biet -> so chu thich
    rungs: int
    rungs_skipped: int


def _ro(path: Path) -> sqlite3.Connection:
    return sqlite3.connect(f"file:{path}?mode=ro", uri=True)


def _ladder_devices(data: str) -> list[tuple[str, int]] | None:
    """Mot khoi ladder: phan dau liet ke loai thiet bi theo thu tu (xen lan ten lenh, hang so K_/H_),
    phan sau la so hieu 'a=<so>' cung thu tu. Tra None neu khong ghep duoc."""
    if "cb{" not in data:
        return []
    toks = data.split("cb{", 1)[0].split(":")[1:]
    i = 0
    while i < len(toks) and toks[i].isdigit():
        i += 1
    types = []
    while i < len(toks):
        t = toks[i]
        if t.startswith(("K_", "H_", "E_", '"')):  # hang so: ten + gia tri
            i += 2
            continue
        if t in LADDER_DEVICES:
            types.append(t)
        i += 1
    nums = re.findall(r"d\{s=#:a=(\d+)", data)
    if len(types) != len(nums):
        return None
    out = []
    for t, n in zip(types, nums):
        # X/Y trong ladder luu theo chu so hien thi doc nhu he 16 (X10 -> 0x10)
        out.append((t, int(f"{int(n):x}", 8) if t in OCTAL else int(n)))
    return out


def read_project(path: str | Path) -> Project:
    with tempfile.TemporaryDirectory() as tmp:
        with zipfile.ZipFile(path) as z:
            names = [n for n in z.namelist() if n.endswith((".db", "UnitConfig.dat"))]
            for n in names:  # ten trong zip dung '\', giai nen phang de khong tao thu muc la
                (Path(tmp) / Path(n.replace("\\", "/")).name).write_bytes(z.read(n))
        d = Path(tmp)
        comments: dict[str, str] = {}
        unknown: dict[int, int] = {}
        for f in d.glob("*_DC.db"):
            with _ro(f) as c:
                q = ("select d.DevCode, d.DevNoHigh, d.DevNoLow, d.BitNo, m.CmtData from DEVICE_DATA d "
                     "join COMMENT_DATA m on m.DeviceSEQ = d.SEQ where coalesce(m.DelFlag, 0) = 0")
                for code, hi, lo, bit, text in c.execute(q):
                    dev = DEV_CODES.get(code)
                    if dev is None:
                        unknown[code] = unknown.get(code, 0) + 1
                        continue
                    comments[fmt(dev, (hi << 16) + lo, bit or 0)] = (text or "").strip()
        used: set[str] = set()
        rungs = skipped = 0
        for f in d.glob("*_LDDB.db"):
            with _ro(f) as c:
                for (data,) in c.execute("select data from LadderBlocks"):
                    rungs += 1
                    devs = _ladder_devices(data or "")
                    if devs is None:
                        skipped += 1
                        continue
                    used.update(fmt(t, n) for t, n in devs)
        models: list[str] = []
        cfg = d / "UnitConfig.dat"
        if cfg.exists():
            raw = cfg.read_bytes()
            found = [(m.start(), m.group().decode()) for m in re.finditer(rb"FX5[A-Z]*-\d+[A-Z]+/[A-Z]+", raw)]
            for _, m in sorted(found):
                if m not in models:
                    models.append(m)
            # CPU dung truoc (I/O cua CPU bat dau tu X0), module mo rong giu thu tu trong file
            models.sort(key=lambda m: not re.match(r"FX5[A-Z]*-\d+M", m))
    return Project(comments, used, models, unknown, rungs, skipped)


# ------------------------------------------------------------------ chia I/O theo module

def io_ranges(models: list[str]) -> list[tuple[str, str, int, int]]:
    """[(model, 'X'|'Y', so dau, so cuoi)] theo thu tu lap. FX5U-80M: 40 vao/40 ra.
    FX5-32ER/ET: 16/16; FX5-..EX: toan vao; FX5-..EY: toan ra. Dia chi lam tron len boi so 8."""
    nx = ny = 0
    out = []
    for m in models:
        cpu = re.match(r"FX5[A-Z]*-(\d+)M", m)
        ext = re.match(r"FX5-(\d+)E(X|Y|R|T)", m)
        if cpu:
            i = o = int(cpu.group(1)) // 2
        elif ext:
            n, k = int(ext.group(1)), ext.group(2)
            i, o = {"X": (n, 0), "Y": (0, n)}.get(k, (n // 2, n // 2))
        else:
            continue
        for dev, cnt in (("X", i), ("Y", o)):
            if not cnt:
                continue
            start = nx if dev == "X" else ny
            out.append((m, dev, start, start + cnt - 1))
            size = -(-cnt // 8) * 8
            if dev == "X":
                nx += size
            else:
                ny += size
    return out


def module_of(addr: str, ranges) -> str:
    dev, n = addr[0], int(addr[1:], 8)
    return next((m for m, d, a, b in ranges if d == dev and a <= n <= b), "")


# ------------------------------------------------------------------ so sanh voi ban ve

ANTONYMS = [("OPEN", "CLOSE"), ("UP", "DOWN"), ("ON", "OFF"), ("ORG", "RT"), ("OK", "NG"),
            ("START", "STOP"), ("IN", "OUT"), ("LEFT", "RIGHT"), ("FWD", "BWD")]


def _words(s: str) -> list[str]:
    return re.findall(r"[A-Z0-9]+", norm(s).upper())


def compare_text(prog: str, dwg: str) -> tuple[str, str]:
    """(muc do, giai thich). muc do: KHOP | GAN_GIONG | KHAC | NGUOC_NGHIA."""
    a, b = _words(prog), _words(dwg)
    if a == b:
        return "KHOP", ""
    sa, sb = set(a), set(b)
    for x, y in ANTONYMS:
        if (x in sa - sb and y in sb - sa) or (y in sa - sb and x in sb - sa):
            return "NGUOC_NGHIA", f"{x}/{y}"
    if "SPARE" in sb and "SPARE" not in sa:
        return "KHAC", "ban ve ghi SPARE nhung chuong trinh co chu thich"
    ratio = difflib.SequenceMatcher(None, " ".join(a), " ".join(b)).ratio()
    if ratio >= 0.75 or sa <= sb or sb <= sa:
        return "GAN_GIONG", f"giong {ratio:.0%}"
    return "KHAC", f"giong {ratio:.0%}"


@dataclass
class Diff:
    address: str
    level: str
    program: str
    drawing: str
    sheet: str
    note: str


def compare(project: Project, drawing: dict) -> list[Diff]:
    """drawing: {dia chi: DrawingIO} tu dienkit.pdfio.extract_io."""
    rows: list[Diff] = []
    addrs = {a for a in project.comments if a[0] in OCTAL and "." not in a} | set(drawing)
    for a in sorted(addrs, key=lambda s: (s[0], int(s[1:], 8))):
        prog = project.comments.get(a, "")
        d = drawing.get(a)
        used = "co dung trong chuong trinh" if a in project.used else "KHONG dung trong chuong trinh"
        if d is None:
            rows.append(Diff(a, "THIEU_BAN_VE", prog, "", "", f"co chu thich, {used}"))
            continue
        if not prog:
            rows.append(Diff(a, "THIEU_CHU_THICH", "", d.desc, d.sheet, used))
            continue
        level, note = compare_text(prog, d.desc)
        rows.append(Diff(a, level, prog, d.desc, d.sheet, note))
    # hai dia chi bi dao cho nhau
    by_addr = {r.address: r for r in rows if r.level in ("NGUOC_NGHIA", "KHAC")}
    for r in by_addr.values():
        for o in by_addr.values():
            if o is not r and _words(r.program) == _words(o.drawing) and _words(o.program) == _words(r.drawing):
                r.level, r.note = "DAO_CHO", f"dao voi {o.address}"
    return rows


# ------------------------------------------------------------------ xuat file

def tag_name(comment: str, addr: str, taken: set[str]) -> str:
    s = re.sub(r"[^A-Z0-9]+", "_", norm(comment).upper()).strip("_")[:28]
    if not s or s[0].isdigit():
        s = f"{addr.replace('.', '_')}_{s}".strip("_")
    name = s
    if name.upper() in taken:
        name = f"{s}_{addr.replace('.', '_')}"
    taken.add(name.upper())
    return name


def write_tags_csv(path, project: Project, plc: str) -> int:
    """Bang tag cho dienkit.tags (bo SM/SD he thong)."""
    taken: set[str] = set()
    n = 0
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["plc", "tag", "dia_chi", "kieu", "mo_ta", "hmi", "rw"])
        for addr, text in sorted(project.comments.items(), key=lambda kv: _sort_key(kv[0])):
            dev = re.match(r"[A-Z]+", addr).group()
            if dev in SYSTEM or not text:
                continue
            kieu = "INT" if dev == "D" and "." not in addr else ""
            rw = "R" if dev in ("X", "T") else "RW"
            w.writerow([plc, tag_name(text, addr, taken), addr, kieu, text, "x", rw])
            n += 1
    return n


def write_io_csv(path, project: Project) -> int:
    ranges = io_ranges(project.models)
    rows = [a for a in project.comments if a[0] in OCTAL and "." not in a]
    taken: set[str] = set()
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["dia_chi", "loai", "tag", "mo_ta", "thiet_bi", "dau_day", "module"])
        for a in sorted(rows, key=_sort_key):
            w.writerow([a, "DI" if a[0] == "X" else "DO", tag_name(project.comments[a], a, taken),
                        project.comments[a], "", "", module_of(a, ranges)])
    return len(rows)


def _sort_key(addr: str):
    m = re.match(r"([A-Z]+)(\d+)(?:\.([0-9A-F]))?", addr)
    dev, num, bit = m.group(1), m.group(2), m.group(3)
    return dev, int(num, 8) if dev in OCTAL else int(num), int(bit or "0", 16)


def write_report(path, project: Project, diffs: list[Diff] | None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    red, yellow, green = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FFF2CC", "E2EFDA"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Tong quan"
    ws.append(["Model trong project", ", ".join(project.models) or "(khong doc duoc)"])
    for m, dev, a, b in io_ranges(project.models):
        ws.append([f"  {m}", f"{fmt(dev, a)} - {fmt(dev, b)}"])
    ws.append(["Khoi ladder doc duoc", f"{project.rungs - project.rungs_skipped}/{project.rungs}"])
    ws.append(["Chu thich", len(project.comments)])
    if project.unknown_codes:
        ws.append(["Ma thiet bi chua biet (bo qua)", str(project.unknown_codes)])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60

    ws = wb.create_sheet("Tham chieu")
    ws.append(["Thiet bi", "Chu thich", "Dung trong ladder", "Van de"])
    allx = set(project.comments) | {u for u in project.used if not u.startswith(("SM", "SD"))}
    for a in sorted(allx, key=_sort_key):
        dev = re.match(r"[A-Z]+", a).group()
        if dev in SYSTEM:
            continue
        c = project.comments.get(a, "")
        u = a in project.used
        issue = "dung nhung chua co chu thich" if u and not c else ("co chu thich nhung ladder khong dung (HMI co the dung)" if c and not u else "")
        ws.append([a, c, "x" if u else "", issue])
        if issue:
            for cell in ws[ws.max_row]:
                cell.fill = yellow
    for col, w in zip("ABCD", [10, 40, 10, 34]):
        ws.column_dimensions[col].width = w
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    if diffs is not None:
        ws = wb.create_sheet("So voi ban ve")
        ws.append(["Dia chi", "Muc do", "Chu thich chuong trinh", "Ban ve", "To", "Ghi chu"])
        fill = {"KHOP": None, "GAN_GIONG": green, "KHAC": yellow}
        order = {"DAO_CHO": 0, "NGUOC_NGHIA": 1, "THIEU_BAN_VE": 2, "THIEU_CHU_THICH": 2, "KHAC": 3, "GAN_GIONG": 4, "KHOP": 5}
        for d in sorted(diffs, key=lambda d: (order[d.level], _sort_key(d.address))):
            ws.append([d.address, d.level, d.program, d.drawing, d.sheet, d.note])
            f = fill.get(d.level, red)
            if f:
                for cell in ws[ws.max_row]:
                    cell.fill = f
        for col, w in zip("ABCDEF", [8, 16, 32, 32, 8, 40]):
            ws.column_dimensions[col].width = w
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        wb.move_sheet(ws, -(len(wb.sheetnames) - 1))
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
    wb.save(path)


def main(argv=None) -> int:
    import argparse
    ap = argparse.ArgumentParser(prog="python -m dienkit.gx3",
                                 description="Doc project GX Works3 (.gx3): chu thich, tham chieu cheo, so voi ban ve")
    ap.add_argument("gx3")
    ap.add_argument("--ban-ve", help="Ban ve PDF de so sanh chu thich I/O (dienkit.pdfio)")
    ap.add_argument("-o", "--out", default="gx3_out", help="Thu muc xuat")
    ap.add_argument("--plc", default="PLC1", help="Ten PLC trong bang tag (= ten thiet bi trong EasyBuilder)")
    a = ap.parse_args(argv)

    p = read_project(a.gx3)
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    print(f"Model: {', '.join(p.models) or '?'}")
    print(f"Chu thich: {len(p.comments)}  |  khoi ladder doc duoc {p.rungs - p.rungs_skipped}/{p.rungs}")
    if p.unknown_codes:
        print(f"  Bo qua chu thich cua ma thiet bi chua biet: {p.unknown_codes}")
    n_io = write_io_csv(out / "io_tu_plc.csv", p)
    n_tag = write_tags_csv(out / "tags_tu_plc.csv", p, a.plc)
    print(f"Da ghi {out / 'io_tu_plc.csv'} ({n_io} diem I/O) va {out / 'tags_tu_plc.csv'} ({n_tag} tag)")

    diffs = None
    if a.ban_ve:
        from .pdfio import extract_io
        diffs = compare(p, extract_io(a.ban_ve))
        bad = [d for d in diffs if d.level not in ("KHOP", "GAN_GIONG")]
        print(f"So voi ban ve: {sum(d.level == 'KHOP' for d in diffs)} khop, "
              f"{sum(d.level == 'GAN_GIONG' for d in diffs)} gan giong, {len(bad)} can xem:")
        for d in bad:
            print(f"  {d.address:5} {d.level:15} PLC: {d.program!r:32} ban ve {d.sheet}: {d.drawing!r}  {d.note}")
    write_report(out / "bao_cao_gx3.xlsx", p, diffs)
    print(f"Bao cao: {out / 'bao_cao_gx3.xlsx'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
